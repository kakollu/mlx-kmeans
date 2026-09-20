"""Local CSV/TSV/XLSX workflow, shared by the command line and notebooks."""
import csv
import datetime
import json
from pathlib import Path
import time


def dependencies():
    try:
        import pandas as pd
        import psutil
    except ImportError as e:
        raise ValueError('Install file support: python -m pip install ".[files]"') from e
    return pd, psutil


def batches(path, sheet=None):
    pd, _ = dependencies()
    path = Path(path).expanduser()
    if not path.is_file():
        raise ValueError(f'Local file not found: {path}')
    ext = path.suffix.lower()
    if ext in ['.csv', '.tsv']:
        sep = '\t' if ext == '.tsv' else ','
        with path.open(encoding='utf-8-sig', newline='') as f:
            header = next(csv.reader(f, delimiter=sep), [])
        check_header(header)
        with pd.read_csv(path, sep=sep, dtype=str, keep_default_na=False, encoding='utf-8-sig',
                         chunksize=10_000) as reader:
            yield from reader
    elif ext == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise ValueError('XLSX support needs openpyxl: install ".[files]"') from e
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError(f'Sheet {sheet!r} not found. Available: {", ".join(book.sheetnames)}')
            ws = book[sheet] if sheet else book.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            raw = next(rows, ())
            header = [str(v) if v is not None else '' for v in raw]
            check_header(header)
            chunk = []
            for row in rows:
                chunk.append(['' if v is None else str(v) for v in row])
                if len(chunk) == 10_000:
                    yield pd.DataFrame(chunk, columns=header)
                    chunk = []
            if chunk:
                yield pd.DataFrame(chunk, columns=header)
        finally:
            book.close()
    else:
        raise ValueError('Use a .csv, .tsv, or .xlsx file. Export older Excel files as CSV or XLSX.')


def check_header(header):
    if not header or any(not str(h).strip() for h in header) or len(header) != len(set(header)):
        raise ValueError('The first row must have unique, nonempty column names.')


def inspect_file(path, sheet=None):
    iterator = batches(path, sheet)
    try:
        frame = next(iterator, None)
        if frame is None:
            raise ValueError('The file has no data rows.')
        return frame.head(5)
    finally:
        iterator.close()


def read_table(path, sheet=None):
    pd, psutil = dependencies()
    chunks, used, rows = [], 0, 0
    # Installed RAM sets the budget; current availability is a secondary guard against a busy Mac.
    installed = psutil.virtual_memory().total
    budget = int(installed * .35)
    for chunk in batches(path, sheet):
        used += int(chunk.memory_usage(index=True, deep=True).sum())
        rows += len(chunk)
        estimate = 256*2**20 + used*4 + rows*len(chunk.columns)*32
        if estimate > budget or estimate > psutil.virtual_memory().available * .7:
            raise ValueError('This file exceeds the conservative in-memory budget. Nothing was sampled. '
                             'Close heavy apps or export a smaller file/fewer columns and retry.')
        chunks.append(chunk)
    if not chunks:
        raise ValueError('The file has no data rows.')
    return pd.concat(chunks, ignore_index=True)


def cluster_file(path, *, columns, k, output=None, n_init=10, max_iter=300, tol=1e-4,
                 seed=0, standardize=True, missing='error', sheet=None):
    """Cluster selected numeric features; save original rows with 1-based cluster labels and a summary.

    No sampling or networking. Missing/nonfinite selected values error by default; missing='drop' retains
    excluded rows in the export with a blank label. Returns paths, a pandas summary, and run metadata.
    """
    import numpy as np
    from .estimator import KMeans
    pd, _ = dependencies()
    for name, value in [('k', k), ('n_init', n_init), ('max_iter', max_iter)]:
        if isinstance(value, bool) or not isinstance(value, (int, np.integer)) or value < 1:
            raise ValueError(f'{name} must be a positive integer')
    if not np.isfinite(tol) or tol < 0:
        raise ValueError('tol must be finite and nonnegative')
    if missing not in ['error', 'drop']:
        raise ValueError('missing must be error or drop')
    if isinstance(columns, str) or not columns or len(set(columns)) != len(columns):
        raise ValueError('Pass a list of unique feature column names, e.g. columns=["spend", "visits"].')
    path = Path(path).expanduser().resolve()
    destination = Path(output).expanduser() if output else path.with_name(
        path.stem+'-clusters-'+datetime.datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
    if destination.exists():
        raise ValueError(f'Output already exists: {destination}. Choose a new folder; nothing overwritten.')
    start = time.perf_counter()
    frame = read_table(path, sheet)
    absent = [c for c in columns if c not in frame.columns]
    if absent:
        raise ValueError(f'Columns not found: {absent}. Available: {list(frame.columns)}')
    numeric = frame[list(columns)].apply(pd.to_numeric, errors='coerce')
    values = numeric.to_numpy(dtype=np.float64)
    valid = np.isfinite(values).all(axis=1)
    excluded = int((~valid).sum())
    if excluded and missing == 'error':
        bad = {c:int((~np.isfinite(values[:,i])).sum()) for i,c in enumerate(columns)
               if (~np.isfinite(values[:,i])).any()}
        raise ValueError(f'{excluded} rows have missing/non-numeric/nonfinite selected values: {bad}. '
                         'Clean the file or explicitly choose --missing drop / missing="drop".')
    kept = values[valid]
    if len(kept) < k:
        raise ValueError(f'Only {len(kept)} usable rows for k={k}.')
    means, std = kept.mean(0), kept.std(0)
    constant = [c for c,s in zip(columns,std) if s == 0]
    if len(constant) == len(columns):
        raise ValueError('All selected features are constant; they cannot separate rows into clusters.')
    center = means if standardize else np.zeros(len(columns))
    scale = np.where(std > 0,std,1) if standardize else np.ones(len(columns))
    with np.errstate(over='ignore',invalid='ignore'):
        X = np.ascontiguousarray((kept-center)/scale,dtype=np.float32)
    if not np.isfinite(X).all():
        raise ValueError('Feature values exceed float32 range; rescale the data before clustering.')
    prepare_s = time.perf_counter()-start
    t = time.perf_counter()
    model = KMeans(n_clusters=k,n_init=n_init,max_iter=max_iter,tol=tol,random_state=seed).fit(X)
    fit_s = time.perf_counter()-t
    labels = model.labels_+1  # R-friendly exported group numbers; not ranked categories
    summary = pd.DataFrame(kept,columns=columns).groupby(labels).mean()
    summary.index.name = 'cluster'
    summary = summary.rename(columns={c:'mean_'+c for c in columns})
    counts = np.bincount(labels,minlength=k+1)[1:]
    summary = summary.reindex(range(1,k+1))
    summary.insert(0,'rows',counts)
    summary.insert(1,'share',counts/len(kept))
    summary = summary.reset_index()
    label_name = 'cluster'
    while label_name in frame.columns:
        label_name = 'kmeans_'+label_name
    exported = pd.Series(pd.NA,index=frame.index,dtype='Int64')
    exported.loc[valid] = labels
    frame[label_name] = exported
    metadata = dict(input=str(path),columns=list(columns),sheet=sheet,k=k,n_init=n_init,max_iter=max_iter,
                    tol=tol,tol_definition='relative inertia change (not R or sklearn center-shift tolerance)',
                    seed=seed,standardize=standardize,scale_ddof=0,missing=missing,
                    input_rows=len(frame),clustered_rows=len(kept),excluded_rows=excluded,
                    label_column=label_name,labels='1-based, arbitrary group IDs, not ranked',
                    constant_features=constant,inertia=model.inertia_,iterations=model.n_iter_,
                    preparation_s=prepare_s,fit_including_labels_s=fit_s,
                    feature_mean=center.tolist(),feature_scale=scale.tolist())
    # Create a new folder only after validation and fitting. Existing files are never replaced.
    destination.mkdir(parents=True,exist_ok=False)
    frame.to_csv(destination/'labeled.csv',index=False)
    summary.to_csv(destination/'cluster_summary.csv',index=False)
    (destination/'run.json').write_text(json.dumps(metadata,indent=2))
    np.savez(destination/'model.npz',centers=model.cluster_centers_,mean=center,scale=scale,
             columns=np.array(columns))
    return dict(output=destination.resolve(),summary=summary,run=metadata)
